import asyncio
import html
import io
import logging
import secrets
from datetime import datetime

from aiogram import Bot, F, Router
from aiogram.dispatcher.event.bases import SkipHandler
from aiogram.exceptions import TelegramBadRequest, TelegramForbiddenError, TelegramRetryAfter
from aiogram.filters import Command, CommandObject, StateFilter
from aiogram.fsm.context import FSMContext
from aiogram.fsm.state import State, StatesGroup
from aiogram.types import BufferedInputFile, CallbackQuery, InputProfilePhotoStatic, Message
from openpyxl import Workbook
from openpyxl.styles import Font

import config
import db
import manual_pay
import materials
import services as S
from db import now
from services import btn, fmt_date, kb

log = logging.getLogger(__name__)
router = Router()
router.message.filter(F.from_user.id.in_(config.ADMIN_IDS))
router.callback_query.filter(F.from_user.id.in_(config.ADMIN_IDS))

_upload_lock = asyncio.Lock()


class Upload(StatesGroup):
    files = State()


class Broadcast(StatesGroup):
    content = State()


class Avatar(StatesGroup):
    photo = State()


HELP = """<b>Админка</b>

<b>Канал</b>
Перешли боту любой пост из канала — бот подключит его
/channel — проверить, всё ли в порядке с каналом
/avatar — поменять аватарку бота

<b>Материалы в боте</b> (не нужны, если материалы в канале)
/addsubject Название — добавить предмет
/subjects — список предметов
/upload — загрузить файлы в предмет
/files id — файлы предмета
/movefile id Тема — перенести файл в другую тему
/delfile id — удалить файл
/renamesubject id Название
/delsubject id — удалить предмет со всеми файлами
/preview id — посмотреть, как увидит подписчик

<b>Подписчики</b>
/paid id|@ник тариф — занести оплату переводом вручную
/stats — статистика
/grant id|@ник дни — выдать доступ
/revoke id|@ник — забрать доступ
/info id|@ник — карточка человека
/export — выгрузка в Excel

<b>Маркетинг</b>
/addpromo КОД ПРОЦЕНТ [ЛИМИТ] [ДНЕЙ] — промокод
/promos — промокоды и их эффективность
/delpromo КОД
/gift тариф [кол-во] — подарочные коды (y1, m1, m3)
/courseend ДД.ММ.ГГГГ — дата конца курса (для приза за 5 друзей)
/broadcast all|active|former|never — рассылка

Ссылки для рекламы:
<code>https://t.me/{bot}?start=src_метка</code> — без скидки, просто считать источник
<code>https://t.me/{bot}?start=p_КОД</code> — с промокодом"""


@router.message(Command("admin", "help"))
async def cmd_admin(m: Message, bot: Bot):
    await m.answer(HELP.format(bot=await S.bot_username(bot)))


@router.message(Command("cancel"))
async def cmd_cancel(m: Message, state: FSMContext):
    await state.clear()
    await m.answer("Ок, отменил.")


async def resolve(arg: str):
    arg = (arg or "").strip()
    if arg.startswith("@"):
        r = await db.fetchone("SELECT user_id FROM users WHERE lower(username)=?", arg[1:].lower())
        return r["user_id"] if r else None
    return int(arg) if arg.isdigit() else None


# ---------- канал ----------

async def channel_problems(bot: Bot, chat_id: int) -> list[str]:
    """Чего не хватает боту в канале. Пустой список — всё хорошо."""
    try:
        me = await bot.get_chat_member(chat_id, bot.id)
    except TelegramBadRequest:
        return ["бот не добавлен в канал"]
    if me.status != "administrator":
        return ["бот в канале, но не админ"]
    out = []
    if not getattr(me, "can_invite_users", False):
        out.append("нет права «Добавлять участников» (нужно, чтобы принимать заявки)")
    if not getattr(me, "can_restrict_members", False):
        out.append("нет права «Блокировать пользователей» (нужно, чтобы убирать тех, кто не продлил)")
    return out


@router.message(StateFilter(None), F.forward_origin.type == "channel")
async def connect_channel(m: Message, bot: Bot):
    chat = m.forward_origin.chat
    problems = await channel_problems(bot, chat.id)
    if problems:
        return await m.answer(
            f"Канал «{html.escape(chat.title or '')}» пока не подключить:\n• " + "\n• ".join(problems) +
            "\n\nОткрой канал → Управление → Администраторы → добавь бота с этими правами "
            "и перешли пост ещё раз.")
    await db.set_setting("channel_id", chat.id)
    config.CHANNEL_ID = chat.id
    link = await S.channel_link(bot)
    await m.answer(
        f"✅ Канал «{html.escape(chat.title or '')}» подключён.\n\n"
        "Теперь «📚 Материалы» пускает в канал только подписчиков, а тех, у кого кончилась подписка, "
        "бот убирает сам.\n\n"
        f"Ссылка для входа (по ней бот принимает только оплативших): {link}\n\n"
        "<b>Сделай в настройках канала:</b>\n"
        "1. Тип канала — <b>частный</b>.\n"
        "2. Удали старые пригласительные ссылки: по ним заходят в обход бота.\n"
        "3. Включи «Запретить копирование» — чтобы материалы не пересылали и не сохраняли.\n\n"
        "Кто уже сидит в канале и платил раньше — выдай им доступ: /grant @ник дни. "
        "Иначе бот о них не знает и не напомнит о продлении.")


@router.message(Command("avatar"))
async def cmd_avatar(m: Message, state: FSMContext):
    await state.set_state(Avatar.photo)
    await m.answer("Пришли картинку для аватарки бота — обычным фото, не файлом. /cancel — отмена.")


@router.message(Avatar.photo, F.photo)
async def avatar_photo(m: Message, bot: Bot, state: FSMContext):
    await state.clear()
    buf = io.BytesIO()
    await bot.download(m.photo[-1], destination=buf)
    try:
        await bot.set_my_profile_photo(InputProfilePhotoStatic(photo=BufferedInputFile(buf.getvalue(), "avatar.jpg")))
    except TelegramBadRequest as e:
        return await m.answer(f"Telegram не принял картинку: {e}. Можно поменять через @BotFather → /setuserpic.")
    await m.answer("✅ Аватарка обновлена. У людей она может обновиться не сразу.")


@router.message(Avatar.photo)
async def avatar_not_photo(m: Message):
    await m.answer("Нужна картинка, отправленная как фото. Или /cancel.")


@router.message(Command("channel"))
async def cmd_channel(m: Message, bot: Bot):
    if not config.CHANNEL_ID:
        return await m.answer("Канал не подключён. Сделай бота админом канала и перешли сюда любой пост из него.")
    try:
        chat = await bot.get_chat(config.CHANNEL_ID)
        title = chat.title
    except TelegramBadRequest:
        title = str(config.CHANNEL_ID)
    problems = await channel_problems(bot, config.CHANNEL_ID)
    if problems:
        return await m.answer(f"⚠️ Канал «{html.escape(title)}»: " + "; ".join(problems))
    joined = await db.scalar("SELECT COUNT(*) FROM events WHERE type='joined_channel'")
    await m.answer(f"✅ Канал «{html.escape(title)}» подключён, права в порядке.\n"
                   f"Ссылка для входа: {await S.channel_link(bot)}\nПринято заявок через бота: {joined}")


# ---------- предметы и файлы ----------

@router.message(Command("addsubject"))
async def cmd_addsubject(m: Message, command: CommandObject):
    name = (command.args or "").strip()
    if not name:
        return await m.answer("Формат: /addsubject Микроэкономика")
    pos = (await db.scalar("SELECT MAX(position) FROM subjects") or 0) + 1
    try:
        cur = await db.execute("INSERT INTO subjects(name,position) VALUES(?,?)", name, pos)
    except Exception:
        return await m.answer("Такой предмет уже есть.")
    await m.answer(f"Добавил «{html.escape(name)}» (id {cur.lastrowid}). Теперь /upload, чтобы залить файлы.")


@router.message(Command("subjects"))
async def cmd_subjects(m: Message):
    rows = await db.fetch(
        "SELECT s.id, s.name, COUNT(m.id) n, COUNT(DISTINCT m.topic) t FROM subjects s "
        "LEFT JOIN materials m ON m.subject_id=s.id GROUP BY s.id ORDER BY s.position, s.name"
    )
    if not rows:
        return await m.answer("Предметов пока нет. /addsubject Название")
    await m.answer("\n".join(f"<code>{r['id']}</code> · {html.escape(r['name'])} — {r['n']} файлов, {r['t']} тем"
                             for r in rows))


@router.message(Command("renamesubject"))
async def cmd_renamesubject(m: Message, command: CommandObject):
    parts = (command.args or "").split(maxsplit=1)
    if len(parts) < 2 or not parts[0].isdigit():
        return await m.answer("Формат: /renamesubject id Новое название")
    await db.execute("UPDATE subjects SET name=? WHERE id=?", parts[1], int(parts[0]))
    await m.answer("Готово.")


@router.message(Command("delsubject"))
async def cmd_delsubject(m: Message, command: CommandObject):
    if not (command.args or "").strip().isdigit():
        return await m.answer("Формат: /delsubject id")
    sid = int(command.args)
    await db.execute("DELETE FROM materials WHERE subject_id=?", sid)
    await db.execute("DELETE FROM subjects WHERE id=?", sid)
    await m.answer("Удалил предмет и его файлы.")


@router.message(Command("upload"))
async def cmd_upload(m: Message):
    subjects = await db.fetch("SELECT id, name FROM subjects ORDER BY position, name")
    if not subjects:
        return await m.answer("Сначала добавь предмет: /addsubject Название")
    await m.answer("Куда грузим?", reply_markup=kb(*[[btn(s["name"], f"up:{s['id']}")] for s in subjects]))


@router.callback_query(F.data.startswith("up:"))
async def cb_upload_subject(cb: CallbackQuery, state: FSMContext):
    sid = int(cb.data[3:])
    subj = await db.fetchone("SELECT * FROM subjects WHERE id=?", sid)
    await state.set_state(Upload.files)
    await state.update_data(subject_id=sid, topic=None)
    await cb.message.edit_text(
        f"Грузим в «{html.escape(subj['name'])}».\n\n"
        "• Подпись к файлу = название темы. В альбоме достаточно подписать первый файл.\n"
        "• Файлы без подписи падают в последнюю тему.\n"
        "• /topic Название — сменить тему для следующих файлов.\n"
        "• /done — закончить."
    )
    await cb.answer()


@router.message(Upload.files, Command("topic"))
async def upload_topic(m: Message, command: CommandObject, state: FSMContext):
    topic = (command.args or "").strip()
    if not topic:
        return await m.answer("Формат: /topic Тема 3. Эластичность")
    await state.update_data(topic=topic)
    await m.answer(f"Дальше кладу в тему «{html.escape(topic)}».")


@router.message(Upload.files, Command("done"))
async def upload_done(m: Message, state: FSMContext):
    data = await state.get_data()
    await state.clear()
    n = await db.scalar("SELECT COUNT(*) FROM materials WHERE subject_id=?", data["subject_id"])
    await m.answer(f"Готово. В предмете сейчас {n} файлов. Проверить: /preview {data['subject_id']}")


def extract_file(m: Message):
    if m.document:
        return m.document.file_id, "document", m.document.file_name
    if m.video:
        return m.video.file_id, "video", m.video.file_name or "Видео"
    if m.audio:
        return m.audio.file_id, "audio", m.audio.file_name or m.audio.title or "Аудио"
    if m.photo:
        return m.photo[-1].file_id, "photo", None
    if m.voice:
        return m.voice.file_id, "voice", "Голосовое"
    return None


@router.message(Upload.files)
async def upload_file(m: Message, state: FSMContext):
    if m.text and m.text.startswith("/"):
        raise SkipHandler()  # другие команды работают и во время загрузки
    f = extract_file(m)
    if not f:
        return await m.answer("Это не файл. Пришли документ, видео, аудио или фото — или /done.")
    # Альбом приходит пачкой апдейтов; лок держит порядок, чтобы подпись первого файла
    # успела стать темой для остальных
    async with _upload_lock:
        data = await state.get_data()
        caption = (m.caption or "").strip().split("\n")[0][:100]
        topic = caption or data.get("topic") or materials.NO_TOPIC
        if caption:
            await state.update_data(topic=caption)
        sid = data["subject_id"]
        pos = (await db.scalar("SELECT MAX(position) FROM materials WHERE subject_id=?", sid) or 0) + 1
        file_id, ftype, name = f
        cur = await db.execute(
            "INSERT INTO materials(subject_id,topic,file_id,file_type,file_name,position,created_at) "
            "VALUES(?,?,?,?,?,?,?)", sid, topic, file_id, ftype, name, pos, now())
    await m.reply(f"✓ #{cur.lastrowid} → {html.escape(topic)}")


@router.message(Command("files"))
async def cmd_files(m: Message, command: CommandObject):
    if not (command.args or "").strip().isdigit():
        return await m.answer("Формат: /files id_предмета")
    mats = await db.fetch("SELECT * FROM materials WHERE subject_id=? ORDER BY position, id", int(command.args))
    if not mats:
        return await m.answer("Файлов нет.")
    out = []
    for topic, files in materials.group_by_topic(mats):
        out.append(f"<b>{html.escape(topic)}</b>")
        out += [f"  <code>{f['id']}</code> {html.escape(f['file_name'] or f['file_type'])}" for f in files]
    text = "\n".join(out)
    for i in range(0, len(text), 4000):
        await m.answer(text[i:i + 4000])


@router.message(Command("movefile"))
async def cmd_movefile(m: Message, command: CommandObject):
    parts = (command.args or "").split(maxsplit=1)
    if len(parts) < 2 or not parts[0].isdigit():
        return await m.answer("Формат: /movefile id Новая тема")
    await db.execute("UPDATE materials SET topic=? WHERE id=?", parts[1], int(parts[0]))
    await m.answer("Перенёс.")


@router.message(Command("delfile"))
async def cmd_delfile(m: Message, command: CommandObject):
    if not (command.args or "").strip().isdigit():
        return await m.answer("Формат: /delfile id")
    await db.execute("DELETE FROM materials WHERE id=?", int(command.args))
    await m.answer("Удалил.")


@router.message(Command("preview"))
async def cmd_preview(m: Message, command: CommandObject, bot: Bot):
    if not (command.args or "").strip().isdigit():
        return await m.answer("Формат: /preview id_предмета")
    subj = await db.fetchone("SELECT * FROM subjects WHERE id=?", int(command.args))
    if not subj:
        return await m.answer("Нет такого предмета.")
    if not await materials.send_subject(bot, m.chat.id, subj):
        await m.answer("В предмете пока нет файлов.")


# ---------- подписчики ----------

@router.message(Command("grant"))
async def cmd_grant(m: Message, command: CommandObject, bot: Bot):
    parts = (command.args or "").split()
    if len(parts) < 2 or not parts[1].isdigit():
        return await m.answer("Формат: /grant 123456789 30  или  /grant @ник 30")
    uid = await resolve(parts[0])
    if not uid:
        return await m.answer("Не нашёл. По @нику ищу только тех, кто уже запускал бота.")
    if not await db.get_user(uid):
        await db.execute("INSERT INTO users(user_id,created_at,source) VALUES(?,?,?)", uid, now(), "manual")
    days = int(parts[1])
    exp = await S.extend(bot, uid, days, default="manual", text=f"🎁 Тебе открыли доступ на {days} дн.")
    await db.event(uid, "grant", str(days))
    await m.answer(f"Выдал. Доступ до {fmt_date(exp)}.")


@router.message(Command("revoke"))
async def cmd_revoke(m: Message, command: CommandObject, bot: Bot):
    uid = await resolve(command.args)
    if not uid:
        return await m.answer("Формат: /revoke 123456789  или  /revoke @ник")
    await S.kick(bot, uid)
    await db.execute("UPDATE subs SET status='expired', autorenew=0 WHERE user_id=?", uid)
    await db.event(uid, "revoked")
    await m.answer("Доступ забрал.")


@router.message(Command("info"))
async def cmd_info(m: Message, command: CommandObject):
    uid = await resolve(command.args)
    u = await db.get_user(uid) if uid else None
    if not u:
        return await m.answer("Не нашёл.")
    s = await db.get_sub(uid)
    paid = await db.scalar("SELECT COALESCE(SUM(amount),0) FROM payments WHERE user_id=? AND status='succeeded'", uid)
    invited = await db.scalar("SELECT COUNT(*) FROM users WHERE referrer_id=?", uid)
    sub = "нет"
    if s:
        state = "активна" if s["status"] == "active" and s["expires_at"] > now() else "истекла"
        sub = f"{state} до {fmt_date(s['expires_at'])} · {S.tariff_title(s['tariff'])} · {s['method']}"
    await m.answer(
        f"<b>{html.escape(S.person(u))}</b> @{u['username'] or '—'} · <code>{uid}</code>\n"
        f"С нами с {fmt_date(u['created_at'])}, источник: {u['source'] or '—'}\n"
        f"Подписка: {sub}\nОплатил всего: {paid} ₽\nПривёл друзей: {invited}"
    )


def month_start() -> int:
    return int(datetime.now(config.TZ).replace(day=1, hour=0, minute=0, second=0, microsecond=0).timestamp())


@router.message(Command("stats"))
async def cmd_stats(m: Message):
    t, ms = now(), month_start()
    q = db.scalar
    active = await q("SELECT COUNT(*) FROM subs WHERE status='active' AND expires_at>? AND method!='trial'", t)
    trials = await q("SELECT COUNT(*) FROM subs WHERE status='active' AND expires_at>? AND method='trial'", t)
    users = await q("SELECT COUNT(*) FROM users")
    new_users = await q("SELECT COUNT(*) FROM users WHERE created_at>=?", ms)
    revenue = await q("SELECT COALESCE(SUM(amount),0) FROM payments WHERE status='succeeded' AND currency='RUB' AND paid_at>=?", ms)
    test_rev = await q("SELECT COALESCE(SUM(amount),0) FROM payments WHERE status='succeeded' AND provider='test' AND paid_at>=?", ms)
    new_paid = await q(
        "SELECT COUNT(*) FROM (SELECT user_id, MIN(paid_at) f FROM payments WHERE status='succeeded' "
        "AND kind!='gift' GROUP BY user_id) WHERE f>=?", ms)
    churn = await q("SELECT COUNT(*) FROM events WHERE type='expired' AND note!='trial' AND at>=?", ms)
    trial_end = await q("SELECT COUNT(*) FROM events WHERE type='expired' AND note='trial' AND at>=?", ms)
    trials_month = await q("SELECT COUNT(*) FROM events WHERE type='trial' AND at>=?", ms)
    gifts = await q("SELECT COUNT(*) FROM gifts WHERE created_at>=?", ms)
    by_bank = await db.fetch(
        "SELECT bank, COUNT(*) n, SUM(amount) rub FROM payments WHERE status='succeeded' AND provider='manual' "
        "AND bank IS NOT NULL AND paid_at>=? GROUP BY bank", ms)
    top = await db.fetch(
        "SELECT s.name, COUNT(*) n FROM opens o JOIN subjects s ON s.id=o.subject_id "
        "WHERE o.at>=? GROUP BY s.id ORDER BY n DESC LIMIT 5", ms)
    sources = await db.fetch(
        "SELECT COALESCE(u.source,'—') src, COUNT(*) users, "
        "SUM(EXISTS(SELECT 1 FROM payments p WHERE p.user_id=u.user_id AND p.status='succeeded')) payers, "
        "COALESCE(SUM((SELECT SUM(amount) FROM payments p WHERE p.user_id=u.user_id AND p.status='succeeded')),0) rub "
        "FROM users u GROUP BY src ORDER BY users DESC LIMIT 15")

    lines = [
        "<b>Сейчас</b>",
        f"Активных подписок: {active}",
        *([f"На пробном: {trials}"] if config.TRIAL_DAYS or trials else []),
        f"Всего в боте: {users}",
        "",
        "<b>Этот месяц</b>",
        f"Выручка: {revenue} ₽" + (f" (из них тестовых {test_rev} ₽)" if test_rev else ""),
        f"Новых в боте: {new_users}",
        f"Впервые оплатили: {new_paid}",
        f"Не продлили: {churn}",
        *([f"Взяли пробный: {trials_month}, пробный закончился: {trial_end}"]
          if config.TRIAL_DAYS or trials_month or trial_end else []),
        f"Куплено подарков: {gifts}",
    ]
    if by_bank:
        lines += ["", "<b>По банкам</b>"] + [f"{html.escape(r['bank'])}: {r['n']} опл. · {r['rub']} ₽" for r in by_bank]
    if top:
        lines += ["", "<b>Что открывают</b>"] + [f"{html.escape(r['name'])} — {r['n']}" for r in top]
    if sources:
        lines += ["", "<b>Источники</b> (люди / оплатили / ₽)"]
        lines += [f"{html.escape(r['src'])}: {r['users']} / {r['payers']} / {r['rub']}" for r in sources]
    await m.answer("\n".join(lines))


METHODS = {"manual": "перевод СБП", "test": "тестовая", "trial": "пробный", "gift": "подарок",
           "bonus": "бонус за друга", "yk": "ЮKassa", "stars": "звёзды"}
PAY_STATUS = {"succeeded": "оплачен", "canceled": "отменён", "pending": "не оплачен",
              "expired": "не оплачен", "rejected": "отклонён"}
PAY_KIND = {"sub": "подписка", "renew": "продление", "gift": "подарок"}


@router.message(Command("export"))
async def cmd_export(m: Message):
    def ts(v):
        return datetime.fromtimestamp(v, config.TZ).replace(tzinfo=None) if v else None

    names = {r["user_id"]: (r["fio"] or (f"@{r['username']}" if r["username"] else r["full_name"]))
             for r in await db.fetch("SELECT user_id, username, full_name, fio FROM users")}

    wb = Workbook()
    ws = wb.active
    ws.title = "Подписчики"
    head = ["ФИО", "Группа", "Имя в Telegram", "Ник", "ID", "Статус", "Тариф", "Доступ до", "Оплатил, ₽", "Банк",
            "Откуда пришёл", "Пригласил", "В боте с"]
    if config.CHANNEL_ID:
        head.insert(6, "В канале")
    ws.append(head)
    rows = await db.fetch(
        "SELECT u.*, s.status, s.tariff, s.method, s.expires_at, "
        "(SELECT COALESCE(SUM(amount),0) FROM payments p WHERE p.user_id=u.user_id AND p.status='succeeded') paid "
        "FROM users u LEFT JOIN subs s ON s.user_id=u.user_id "
        "ORDER BY (s.status='active' AND s.expires_at>?) DESC, s.expires_at DESC, u.created_at DESC", now())
    for r in rows:
        if not r["status"]:
            status = "не подписывался"
        elif r["status"] == "active" and r["expires_at"] > now():
            status = "пробный" if r["method"] == "trial" else "активна"
        else:
            status = "истекла"
        source = {"ref": "по приглашению"}.get(r["source"], r["source"])
        if source and source.startswith("src_"):
            source = source[4:]
        row = [r["fio"], r["study_group"], r["full_name"], f"@{r['username']}" if r["username"] else None,
               r["user_id"], status,
               S.tariff_title(r["tariff"]) if r["tariff"] else None, ts(r["expires_at"]), r["paid"],
               r["pay_bank"], source, names.get(r["referrer_id"]), ts(r["created_at"])]
        if config.CHANNEL_ID:
            row.insert(6, "да" if await S.in_channel(m.bot, r["user_id"]) else "нет")
            await asyncio.sleep(0.04)  # не упираться в лимиты Telegram
        ws.append(row)

    ws2 = wb.create_sheet("Платежи")
    ws2.append(["Дата", "Кто", "ID", "Сумма, ₽", "Тариф", "Что", "Способ", "Банк", "Промокод", "Статус"])
    # первая оплата человека — «подписка», следующие — «продление»
    first = {r[0] for r in await db.fetch(
        "SELECT MIN(id) FROM payments WHERE status='succeeded' AND kind!='gift' GROUP BY user_id")}
    for p in await db.fetch("SELECT * FROM payments WHERE status IN ('succeeded','canceled') "
                            "ORDER BY COALESCE(paid_at, created_at) DESC"):
        ws2.append([ts(p["paid_at"] or p["created_at"]), names.get(p["user_id"]), p["user_id"], p["amount"],
                    S.tariff_title(p["tariff"]),
                    "продление" if p["kind"] == "sub" and p["status"] == "succeeded" and p["id"] not in first
                    else PAY_KIND.get(p["kind"], p["kind"]),
                    METHODS.get(p["provider"], p["provider"]), p["bank"], p["promo"],
                    PAY_STATUS.get(p["status"], p["status"])])

    for sheet in (ws, ws2):
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for c in sheet[1]:
            c.font = Font(bold=True)
        for col in sheet.columns:
            sheet.column_dimensions[col[0].column_letter].width = max(
                10, min(32, max(len(str(c.value or "")) for c in col) + 3))
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value, datetime):
                    c.number_format = "DD.MM.YYYY"

    buf = io.BytesIO()
    wb.save(buf)
    name = f"subscribers_{datetime.now(config.TZ):%Y-%m-%d}.xlsx"
    await m.answer_document(BufferedInputFile(buf.getvalue(), filename=name))


# ---------- промокоды и подарки ----------

@router.message(Command("addpromo"))
async def cmd_addpromo(m: Message, command: CommandObject, bot: Bot):
    p = (command.args or "").split()
    if len(p) < 2 or not p[1].isdigit() or not 1 <= int(p[1]) <= 99:
        return await m.answer("Формат: /addpromo VK20 20 [лимит] [дней]\nНапример: /addpromo VK20 20 100 14")
    code, pct = p[0].upper(), int(p[1])
    uses = int(p[2]) if len(p) > 2 and p[2].isdigit() else None
    exp = now() + int(p[3]) * 86400 if len(p) > 3 and p[3].isdigit() else None
    await db.execute("INSERT OR REPLACE INTO promos(code,percent,max_uses,expires_at,created_at) VALUES(?,?,?,?,?)",
                     code, pct, uses, exp, now())
    link = f"https://t.me/{await S.bot_username(bot)}?start=p_{code}"
    await m.answer(f"Промокод <b>{code}</b> −{pct}% создан.\nСсылка для рекламы:\n<code>{link}</code>")


@router.message(Command("promos"))
async def cmd_promos(m: Message):
    rows = await db.fetch(
        "SELECT p.*, (SELECT COUNT(*) FROM users u WHERE u.source=p.code) came, "
        "(SELECT COALESCE(SUM(amount),0) FROM payments x WHERE x.promo=p.code AND x.status='succeeded') rub "
        "FROM promos p WHERE p.user_id IS NULL ORDER BY p.created_at DESC")
    if not rows:
        return await m.answer("Промокодов нет.")
    out = []
    for r in rows:
        lim = f"/{r['max_uses']}" if r["max_uses"] else ""
        exp = f", до {fmt_date(r['expires_at'])}" if r["expires_at"] else ""
        out.append(f"<b>{r['code']}</b> −{r['percent']}% · исп. {r['uses']}{lim}{exp} · "
                   f"пришло {r['came']} · {r['rub']} ₽")
    await m.answer("\n".join(out))


@router.message(Command("delpromo"))
async def cmd_delpromo(m: Message, command: CommandObject):
    await db.execute("DELETE FROM promos WHERE code=?", (command.args or "").strip().upper())
    await m.answer("Удалил.")


@router.message(Command("courseend"))
async def cmd_courseend(m: Message, command: CommandObject):
    arg = (command.args or "").strip()
    if not arg:
        return await m.answer(f"Конец курса сейчас: {fmt_date(await S.course_end_ts())}.\n"
                              "Поменять: /courseend 30.06.2027")
    try:
        datetime.strptime(arg, "%d.%m.%Y")
    except ValueError:
        return await m.answer("Формат: /courseend 30.06.2027")
    await db.set_setting("course_end", arg)
    await m.answer(f"Готово, конец курса — {arg}. Приз за {config.REF_BIG_EVERY} друзей даёт доступ до этой даты.")


@router.message(Command("gift"))
async def cmd_gift(m: Message, command: CommandObject, bot: Bot):
    p = (command.args or "").split()
    if not p or p[0] not in config.TARIFFS:
        return await m.answer(f"Формат: /gift m1 [кол-во]. Тарифы: {', '.join(k for k, _ in config.sale_tariffs())}")
    n = min(int(p[1]), 50) if len(p) > 1 and p[1].isdigit() else 1
    bot_name = await S.bot_username(bot)
    links = []
    for _ in range(n):
        code = "GIFT-" + secrets.token_hex(4).upper()
        await db.execute("INSERT INTO gifts(code,tariff,buyer_id,created_at) VALUES(?,?,?,?)",
                         code, p[0], m.from_user.id, now())
        links.append(f"https://t.me/{bot_name}?start=g_{code}")
    await m.answer("Подарочные ссылки (каждая на одного человека):\n" + "\n".join(links))


# ---------- оплата переводом по СБП ----------

@router.callback_query(F.data.startswith("mp_undo:"))
async def cb_manual_undo(cb: CallbackQuery, bot: Bot):
    if not await manual_pay.undo(bot, int(cb.data.split(":")[1])):
        return await cb.answer("Уже отменено", show_alert=True)
    try:
        await cb.message.edit_text(f"{cb.message.html_text}\n\n<b>↩️ Отменено</b>", reply_markup=None)
    except TelegramBadRequest:
        pass
    await cb.answer("Отменил")


@router.message(Command("paid"))
async def cmd_paid(m: Message, command: CommandObject, bot: Bot):
    """Человек перевёл, но не нажал кнопку — заносим сами."""
    p = (command.args or "").split()
    if len(p) < 2 or p[1] not in config.TARIFFS:
        names = ", ".join(f"{k} — {t['title']}" for k, t in config.sale_tariffs())
        return await m.answer(f"Формат: /paid @ник m1  или  /paid 123456789 y1\nТарифы: {names}")
    uid = await resolve(p[0])
    if not uid or not await db.get_user(uid):
        return await m.answer("Не нашёл. Человек должен хотя бы раз запустить бота.")
    pid = await S.create_payment(uid, "manual", p[1], "sub", config.TARIFFS[p[1]]["rub"])
    await db.execute("UPDATE payments SET bank=? WHERE id=?", await manual_pay.bank_for(uid), pid)
    await manual_pay.mark_paid(bot, pid, by_admin=True)


# ---------- рассылка ----------

AUDIENCES = {
    "all": ("всем", "SELECT user_id FROM users"),
    "active": ("с активной подпиской", "SELECT user_id FROM subs WHERE status='active' AND expires_at>{t}"),
    "former": ("бывшим подписчикам", "SELECT user_id FROM subs WHERE NOT (status='active' AND expires_at>{t})"),
    "never": ("тем, кто ни разу не подписывался", "SELECT user_id FROM users WHERE user_id NOT IN (SELECT user_id FROM subs)"),
}


@router.message(Command("broadcast"))
async def cmd_broadcast(m: Message, command: CommandObject, state: FSMContext):
    aud = (command.args or "").strip()
    if aud not in AUDIENCES:
        return await m.answer("Формат: /broadcast all | active | former | never")
    await state.set_state(Broadcast.content)
    await state.update_data(aud=aud)
    await m.answer(f"Рассылка {AUDIENCES[aud][0]}. Пришли сообщение (текст, фото, что угодно). /cancel — отмена.")


@router.message(Broadcast.content)
async def broadcast_content(m: Message, state: FSMContext):
    data = await state.get_data()
    ids = [r[0] for r in await db.fetch(AUDIENCES[data["aud"]][1].format(t=now()))]
    await state.update_data(chat=m.chat.id, msg=m.message_id, ids=ids)
    await m.answer(f"Отправить это {len(ids)} людям?",
                   reply_markup=kb([btn("Отправить", "bc:go"), btn("Отмена", "bc:no")]))


@router.callback_query(F.data.startswith("bc:"))
async def cb_broadcast(cb: CallbackQuery, state: FSMContext, bot: Bot):
    data = await state.get_data()
    await state.clear()
    if cb.data == "bc:no" or "ids" not in data:
        await cb.message.edit_text("Отменил.")
        return await cb.answer()
    await cb.message.edit_text(f"Рассылаю {len(data['ids'])} людям…")
    asyncio.create_task(run_broadcast(bot, data["ids"], data["chat"], data["msg"], cb.from_user.id))
    await cb.answer()


async def run_broadcast(bot: Bot, ids, from_chat, msg_id, admin_id):
    ok = fail = 0
    for uid in ids:
        for _ in range(3):
            try:
                await bot.copy_message(uid, from_chat, msg_id)
                ok += 1
                break
            except TelegramRetryAfter as e:
                await asyncio.sleep(e.retry_after + 1)
            except (TelegramForbiddenError, TelegramBadRequest):
                fail += 1
                break
        await asyncio.sleep(0.05)
    await S.safe_send(bot, admin_id, f"Рассылка закончена: доставлено {ok}, не дошло {fail}.")


# ---------- ответы на вопросы ----------
# Админ отвечает реплаем на вопрос в чате поддержки — ответ уходит человеку

@router.message(F.chat.id == config.SUPPORT_CHAT_ID, F.reply_to_message)
async def support_reply(m: Message, bot: Bot):
    if m.text and m.text.startswith("/"):
        raise SkipHandler()
    row = await db.fetchone("SELECT user_id FROM support WHERE admin_msg_id=?", m.reply_to_message.message_id)
    if not row:
        raise SkipHandler()
    try:
        await bot.send_message(row["user_id"], "💬 Ответ админа:")
        await m.copy_to(row["user_id"])
        await m.reply("Отправил ✅")
    except (TelegramForbiddenError, TelegramBadRequest):
        await m.reply("Не доставлено: человек заблокировал бота.")
